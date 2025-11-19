from __future__ import annotations

import json
import sqlite3
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from PySide6.QtCore import QAbstractTableModel, QModelIndex, Qt


class TransactionModel(QAbstractTableModel):
    headers = [
        "ФИО",
        "Штрихкод",
        "Предмет",
        "Ряд",
        "Сектор",
        "Дата",
        "Тип",
    ]

    def __init__(self, data: List[Tuple[Any, ...]]):
        super().__init__()
        self._data = data

    def rowCount(self, parent: QModelIndex | None = None) -> int:
        return len(self._data)

    def columnCount(self, parent: QModelIndex | None = None) -> int:
        return len(self.headers)

    def data(self, index: QModelIndex, role: int = Qt.DisplayRole):
        if not index.isValid() or role not in (Qt.DisplayRole, Qt.EditRole):
            return None
        row = self._data[index.row()]
        mapping = (row[1], row[2], row[3], row[4], row[5], row[6], row[7])
        return mapping[index.column()]

    def headerData(self, section: int, orientation: Qt.Orientation, role: int = Qt.DisplayRole):
        if role != Qt.DisplayRole:
            return None
        if orientation == Qt.Horizontal:
            return self.headers[section]
        return str(section + 1)


class DbManager:
    def __init__(self, db_path: Path):
        self.db_path = db_path
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self._conn = sqlite3.connect(self.db_path)
        self._conn.row_factory = sqlite3.Row
        self._init_schema()

    def _init_schema(self) -> None:
        cur = self._conn.cursor()
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS transactions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                fio TEXT NOT NULL,
                barcode TEXT NOT NULL,
                item TEXT NOT NULL,
                row_no INTEGER NOT NULL,
                sector INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                type TEXT NOT NULL
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS inventory (
                barcode TEXT PRIMARY KEY,
                fio TEXT,
                item TEXT,
                row_no INTEGER NOT NULL,
                sector INTEGER NOT NULL,
                updated_at TEXT NOT NULL
            )
            """
        )
        cur.execute(
            """
            CREATE TABLE IF NOT EXISTS state (
                key TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
            """
        )
        self._conn.commit()

    def log_transaction(
        self,
        fio: str,
        barcode: str,
        item: str,
        row_no: int,
        sector: int,
        action_type: str,
        created_at: str,
    ) -> None:
        cur = self._conn.cursor()
        cur.execute(
            "INSERT INTO transactions (fio, barcode, item, row_no, sector, created_at, type) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (fio, barcode, item, row_no, sector, created_at, action_type),
        )
        self._conn.commit()

    def set_inventory(self, barcode: str, fio: str, item: str, row_no: int, sector: int, ts: str) -> None:
        cur = self._conn.cursor()
        cur.execute(
            "REPLACE INTO inventory (barcode, fio, item, row_no, sector, updated_at) VALUES (?, ?, ?, ?, ?, ?)",
            (barcode, fio, item, row_no, sector, ts),
        )
        self._conn.commit()

    def remove_inventory(self, barcode: str) -> Optional[sqlite3.Row]:
        cur = self._conn.cursor()
        cur.execute("SELECT * FROM inventory WHERE barcode = ?", (barcode,))
        row = cur.fetchone()
        if row:
            cur.execute("DELETE FROM inventory WHERE barcode = ?", (barcode,))
            self._conn.commit()
        return row

    def get_inventory(self, barcode: str) -> Optional[sqlite3.Row]:
        cur = self._conn.cursor()
        cur.execute("SELECT * FROM inventory WHERE barcode = ?", (barcode,))
        return cur.fetchone()

    def fetch_transactions(self, limit: int = 200) -> List[Tuple[Any, ...]]:
        cur = self._conn.cursor()
        cur.execute(
            "SELECT * FROM transactions ORDER BY datetime(created_at) DESC LIMIT ?",
            (limit,),
        )
        return cur.fetchall()

    def create_model(self) -> TransactionModel:
        return TransactionModel(self.fetch_transactions())

    def export_excel(self, path: Path) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["ФИО", "Штрихкод", "Предмет", "Ряд", "Сектор", "Дата", "Тип"])
        for row in self.fetch_transactions(limit=10000):
            ws.append([row[1], row[2], row[3], row[4], row[5], row[6], row[7]])
        wb.save(path)

    def import_excel(self, path: Path) -> None:
        from openpyxl import load_workbook

        wb = load_workbook(path)
        ws = wb.active
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        for fio, barcode, item, row_no, sector, created_at, action_type in rows:
            self.log_transaction(
                fio=fio,
                barcode=barcode,
                item=item,
                row_no=int(row_no),
                sector=int(sector),
                action_type=action_type,
                created_at=created_at,
            )

    def export_csv(self, path: Path) -> None:
        import csv

        with path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.writer(fh)
            writer.writerow(["ФИО", "Штрихкод", "Предмет", "Ряд", "Сектор", "Дата", "Тип"])
            for row in self.fetch_transactions(limit=10000):
                writer.writerow([row[1], row[2], row[3], row[4], row[5], row[6], row[7]])

    def import_csv(self, path: Path) -> None:
        import csv

        with path.open("r", newline="", encoding="utf-8") as fh:
            reader = csv.DictReader(fh)
            for row in reader:
                self.log_transaction(
                    fio=row["ФИО"],
                    barcode=row["Штрихкод"],
                    item=row["Предмет"],
                    row_no=int(row["Ряд"]),
                    sector=int(row["Сектор"]),
                    action_type=row["Тип"],
                    created_at=row["Дата"],
                )

    def set_state(self, key: str, value: Dict[str, Any]) -> None:
        cur = self._conn.cursor()
        cur.execute(
            "REPLACE INTO state (key, value) VALUES (?, ?)",
            (key, json.dumps(value)),
        )
        self._conn.commit()

    def get_state(self, key: str) -> Optional[Dict[str, Any]]:
        cur = self._conn.cursor()
        cur.execute("SELECT value FROM state WHERE key = ?", (key,))
        row = cur.fetchone()
        if row:
            return json.loads(row[0])
        return None

    def close(self) -> None:
        self._conn.close()
